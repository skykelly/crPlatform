#!/usr/bin/env python3
"""Official-statistics ingestion, validation and static-site export."""

import argparse
import datetime as dt
import json
import os
import re
import sqlite3
import sys
import urllib.parse
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / "data" / "factbase.sqlite"
SITE_DATA = ROOT / "sites" / "korea-life-fact-base" / "dist" / "data" / "dashboard.json"
WORKBOOK = ROOT / "outputs" / "01a1263b-a1e2-7bf1-92f4-1494b204937a" / "korea_customer_life_fact_dataset_phase1.xlsx"
API = "https://kosis.kr/openapi/Param/statisticsParameterData.do"
TABLES = {
    "DT_1PL1502": ("Demographic", "1인가구", "Actual"),
    "DT_1BZ0506": ("Demographic", "1인가구·고령가구", "Official Projection"),
    "DT_1JC1517": ("Demographic", "고령가구·부부가구", "Actual"),
    "DT_1JC1516": ("Demographic", "가구규모·세대구성", "Actual"),
    "DT_1IN1502": ("Demographic", "인구·가구·주택", "Actual"),
    "DT_1B8000G": ("Life Event", "결혼·출산", "Actual"),
    "DT_1JU1501": ("Home", "주거 형태", "Actual"),
    "DT_1JC1521": ("Home", "평형", "Actual"),
}
FIELDS = ["fact_id", "driver", "topic", "dataset_id", "actual_projection", "year", "geography_code", "geography_name", "dimension_2_code", "dimension_2_name", "dimension_3_code", "dimension_3_name", "measure_code", "measure_name", "value_raw", "unit", "source_url", "retrieved_at"]


def connect(path=DB):
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS catalog (fact_id TEXT PRIMARY KEY, record_json TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS observations (
      fact_id TEXT NOT NULL, driver TEXT NOT NULL, topic TEXT NOT NULL,
      dataset_id TEXT NOT NULL, actual_projection TEXT NOT NULL, year INTEGER NOT NULL,
      geography_code TEXT NOT NULL, geography_name TEXT NOT NULL,
      dimension_2_code TEXT NOT NULL DEFAULT '', dimension_2_name TEXT NOT NULL DEFAULT '',
      dimension_3_code TEXT NOT NULL DEFAULT '', dimension_3_name TEXT NOT NULL DEFAULT '',
      measure_code TEXT NOT NULL, measure_name TEXT NOT NULL, value_raw TEXT NOT NULL,
      unit TEXT NOT NULL, source_url TEXT NOT NULL, retrieved_at TEXT NOT NULL,
      PRIMARY KEY (fact_id,dataset_id,actual_projection,year,geography_code,
                   dimension_2_code,dimension_3_code,measure_code)
    );
    CREATE TABLE IF NOT EXISTS revisions (
      changed_at TEXT NOT NULL, dataset_id TEXT NOT NULL, year INTEGER NOT NULL,
      series_key TEXT NOT NULL, old_value TEXT NOT NULL, new_value TEXT NOT NULL
    );
    CREATE INDEX IF NOT EXISTS observation_lookup ON observations(dataset_id,year,geography_code,measure_code);
    """)
    return conn


def source_url(table):
    return f"https://kosis.kr/statHtml/statHtml.do?orgId=101&tblId={table}&conn_path=I2"


def fact_for(table, measure, name):
    if table == "DT_1PL1502": return "D01"
    if table == "DT_1BZ0506": return "D02"
    if table == "DT_1JC1517":
        return "D06" if measure in ("T210", "T310") else "D01" if measure == "T600" else "D03" if measure == "T100" else "D03/D06 context"
    if table == "DT_1JC1516": return "D08" if "평균" in name else "D09"
    if table == "DT_1IN1502": return "D08 context"
    if table == "DT_1B8000G":
        return "L04" if measure in ("10", "11", "12") else "L03" if measure in ("20", "21", "23", "24") else "L03/L04 context"
    return "H01" if table == "DT_1JU1501" else "H02"


def normalized_value(value):
    if value is None: return ""
    if isinstance(value, (int, float)): return str(value)
    return str(value).strip().replace(",", "")


def insert_row(conn, row, revisions=False):
    row = {key: row[key] for key in FIELDS}
    row["value_raw"] = normalized_value(row["value_raw"])
    key_fields = ("fact_id", "dataset_id", "actual_projection", "year", "geography_code", "dimension_2_code", "dimension_3_code", "measure_code")
    key = tuple(row[k] for k in key_fields)
    where = " AND ".join(f"{k}=?" for k in key_fields)
    old = conn.execute(f"SELECT value_raw FROM observations WHERE {where}", key).fetchone()
    if old and old[0] == row["value_raw"]:
        return False
    if old and old[0] != row["value_raw"]:
        if row["actual_projection"] == "Official Projection":
            raise ValueError(f"Projection changed; review vintage before ingest: {key}")
        if revisions:
            conn.execute("INSERT INTO revisions VALUES (?,?,?,?,?,?)", (row["retrieved_at"], row["dataset_id"], row["year"], json.dumps(key, ensure_ascii=False), old[0], row["value_raw"]))
    columns = ",".join(FIELDS)
    placeholders = ",".join("?" for _ in FIELDS)
    updates = ",".join(f"{k}=excluded.{k}" for k in FIELDS if k not in key_fields)
    conn.execute(f"INSERT INTO observations ({columns}) VALUES ({placeholders}) ON CONFLICT DO UPDATE SET {updates}", tuple(row[k] for k in FIELDS))
    return old is None or old[0] != row["value_raw"]


def bootstrap():
    from openpyxl import load_workbook
    if DB.exists(): raise SystemExit(f"Database exists: {DB}; bootstrap will not overwrite it")
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    conn = connect()
    try:
        with conn:
            catalog = wb["Core Catalog"].iter_rows(min_row=4, max_col=22, values_only=True)
            headers = next(catalog)
            for cells in catalog:
                if cells[0]: conn.execute("INSERT INTO catalog VALUES (?,?)", (str(cells[0]), json.dumps(dict(zip(headers, cells)), ensure_ascii=False)))
            rows = wb["Fact Data Long"].iter_rows(values_only=True)
            next(rows)
            for cells in rows:
                if not cells[0]: continue
                row = dict(zip(FIELDS, cells[1:]))
                for field in ("dimension_2_code", "dimension_2_name", "dimension_3_code", "dimension_3_name"):
                    row[field] = row[field] or ""
                insert_row(conn, row)
        check(conn)
    except Exception:
        conn.close()
        DB.unlink(missing_ok=True)
        raise
    finally:
        conn.close()


def fetch(table, api_key, current_year):
    projection = TABLES[table][2] == "Official Projection"
    params = {
        "method": "getList", "apiKey": api_key, "orgId": "101", "tblId": table,
        "objL1": "00", "objL2": "ALL", "objL3": "ALL", "itmId": "ALL", "prdSe": "Y",
        "startPrdDe": "2022" if projection else str(current_year - 2),
        "endPrdDe": "2052" if projection else str(current_year), "format": "json",
    }
    url = API + "?" + urllib.parse.urlencode(params)
    request = urllib.request.Request(url, headers={"User-Agent": "crPlatform-factbase/1.0"})
    try:
        with urllib.request.urlopen(request, timeout=90) as response:
            body = response.read().decode("utf-8", errors="replace")
    except Exception as exc:
        raise RuntimeError(f"{table}: KOSIS API request failed ({type(exc).__name__})") from None
    try:
        payload = json.loads(body)
    except json.JSONDecodeError:
        # KOSIS documents JSON, but API errors can use JS object notation:
        # {err:"11",errMsg:"..."}. Never print the URL or key.
        error_code = re.search(r'err\s*:\s*"([0-9A-Za-z_-]+)"', body)
        error_message = re.search(r'errMsg\s*:\s*"([^"]{0,200})"', body)
        if error_code:
            message = (error_message.group(1) if error_message else "API error").replace(api_key, "[redacted]")
            raise ValueError(f"{table}: KOSIS error {error_code.group(1)}: {message}") from None
        raise ValueError(f"{table}: KOSIS returned non-JSON data") from None
    if not isinstance(payload, list) or not payload:
        raise ValueError(f"{table}: API returned no data rows (response type {type(payload).__name__})")
    return payload


def from_api(table, item, today):
    if item.get("TBL_ID") != table: raise ValueError(f"Wrong table in API response: {table}")
    if item.get("PRD_SE") not in (None, "Y"): raise ValueError(f"Nonannual data returned: {table}")
    if any(item.get(f"C{i}") for i in range(4, 9)): raise ValueError(f"Unexpected extra dimensions: {table}")
    year = int(item["PRD_DE"])
    if not 2010 <= year <= 2052: raise ValueError(f"Unexpected year: {table}/{year}")
    if item.get("C1") != "00": raise ValueError(f"Non-national geography: {table}/{item.get('C1')}")
    measure = str(item["ITM_ID"])
    name = item.get("ITM_NM") or measure
    driver, topic, mode = TABLES[table]
    return dict(fact_id=fact_for(table, measure, name), driver=driver, topic=topic,
                dataset_id=table, actual_projection=mode, year=year,
                geography_code="00", geography_name=item.get("C1_NM") or "전국",
                dimension_2_code=item.get("C2") or "", dimension_2_name=item.get("C2_NM") or "",
                dimension_3_code=item.get("C3") or "", dimension_3_name=item.get("C3_NM") or "",
                measure_code=measure, measure_name=name, value_raw=normalized_value(item.get("DT")),
                unit=item.get("UNIT_NM") or "지표별", source_url=source_url(table), retrieved_at=today)


def check(conn):
    count = conn.execute("SELECT count(*) FROM observations").fetchone()[0]
    catalog_count = conn.execute("SELECT count(*) FROM catalog").fetchone()[0]
    if count < 80000 or catalog_count != 39: raise ValueError(f"Baseline incomplete: {count} rows, {catalog_count} facts")
    if conn.execute("SELECT count(*) FROM observations WHERE source_url='' OR value_raw='' OR dataset_id='' OR year IS NULL").fetchone()[0]:
        raise ValueError("Required observation fields missing")
    modes = set(r[0] for r in conn.execute("SELECT DISTINCT actual_projection FROM observations"))
    if modes != {"Actual", "Official Projection"}: raise ValueError(f"Invalid status labels: {modes}")
    if conn.execute("SELECT count(*) FROM observations WHERE dataset_id='DT_1BZ0506' AND actual_projection!='Official Projection'").fetchone()[0]:
        raise ValueError("Projection table mixed with actuals")
    if conn.execute("SELECT count(*) FROM observations WHERE dataset_id!='DT_1BZ0506' AND actual_projection!='Actual'").fetchone()[0]:
        raise ValueError("Actual table mixed with projection")
    for table in TABLES:
        if not conn.execute("SELECT 1 FROM observations WHERE dataset_id=? LIMIT 1", (table,)).fetchone():
            raise ValueError(f"Missing source table: {table}")
    return {"observations": count, "catalog": catalog_count, "revisions": conn.execute("SELECT count(*) FROM revisions").fetchone()[0]}


def refresh(api_key):
    if not api_key: raise SystemExit("KOSIS_API_KEY is required. Add it as a GitHub Actions repository secret.")
    today = dt.date.today().isoformat()
    current_year = dt.date.today().year
    responses = {table: fetch(table, api_key, current_year) for table in TABLES}
    staged = []
    for table, items in responses.items():
        converted = [from_api(table, item, today) for item in items]
        identities = [(r["fact_id"], r["year"], r["geography_code"], r["dimension_2_code"], r["dimension_3_code"], r["measure_code"]) for r in converted]
        if len(identities) != len(set(identities)): raise ValueError(f"Duplicate API records: {table}")
        if not converted: raise ValueError(f"Empty API table: {table}")
        staged.extend(converted)
    conn = connect()
    try:
        check(conn)
        changes = 0
        with conn:
            for row in staged: changes += insert_row(conn, row, revisions=True)
            check(conn)
        if changes:
            raw_dir = ROOT / "data" / "raw" / today
            raw_dir.mkdir(parents=True, exist_ok=True)
            for table, items in responses.items():
                (raw_dir / f"{table}.json").write_text(json.dumps(items, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
        print(json.dumps({"retrieved_at": today, "api_rows": len(staged), "changed_observations": changes}, ensure_ascii=False))
    finally:
        conn.close()


def series(conn, table, measure, d2=None, name_part=None):
    query = "SELECT year, value_raw, unit, source_url, retrieved_at FROM observations WHERE dataset_id=? AND geography_code='00' AND measure_code=?"
    args = [table, measure]
    if d2 is not None: query += " AND dimension_2_code=?"; args.append(d2)
    if name_part: query += " AND measure_name LIKE ?"; args.append(f"%{name_part}%")
    query += " ORDER BY year"
    result = []
    for r in conn.execute(query, args):
        try: value = float(r["value_raw"])
        except ValueError: continue
        result.append({"year": r["year"], "value": value, "unit": r["unit"], "source_url": r["source_url"], "retrieved_at": r["retrieved_at"]})
    return result


def build():
    conn = connect()
    try:
        qa = check(conn)
        actual = series(conn, "DT_1PL1502", "T00", "000")
        projection = series(conn, "DT_1BZ0506", "T02", "000")
        births = series(conn, "DT_1B8000G", "10")
        marriages = series(conn, "DT_1B8000G", "20")
        if not actual or not projection or not births or not marriages:
            raise ValueError("A national dashboard series is missing")
        data = {
            "schema_version": 1, "source": "KOSIS", "quality": qa,
            "actual": {"single_person_households": actual, "births": births, "marriages": marriages},
            "official_projection": {"single_person_households": projection},
            "catalog_count": qa["catalog"],
            "coverage": [{"table_id": t, "status": "Official Projection" if t == "DT_1BZ0506" else "Actual",
                          "latest_year": conn.execute("SELECT max(year) FROM observations WHERE dataset_id=?", (t,)).fetchone()[0]}
                         for t in TABLES],
        }
        SITE_DATA.parent.mkdir(parents=True, exist_ok=True)
        SITE_DATA.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
        print(json.dumps({"dashboard": str(SITE_DATA.relative_to(ROOT)), "latest_actual": actual[-1]["year"], **qa}, ensure_ascii=False))
    finally:
        conn.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=("bootstrap", "refresh", "check", "build"))
    args = parser.parse_args()
    if args.command == "bootstrap": bootstrap()
    elif args.command == "refresh": refresh(os.environ.get("KOSIS_API_KEY", ""))
    elif args.command == "build": build()
    else:
        conn = connect()
        try: print(json.dumps(check(conn), ensure_ascii=False))
        finally: conn.close()


if __name__ == "__main__": main()
